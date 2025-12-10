"""
Dosya Isleme Servisi
"""
import io
import logging
from typing import Union

import pandas as pd

logger = logging.getLogger(__name__)


class FileService:
    """Cesitli dosya formatlarini isleyen servis"""

    SUPPORTED_EXTENSIONS = {
        'pdf', 'csv', 'xlsx', 'xls', 'txt', 'docx', 'doc', 'odt', 'json'
    }

    def process(self, content: bytes, extension: str) -> dict:
        """
        Dosyayi uzantisina gore isler.

        Args:
            content: Dosya icerigi (bytes)
            extension: Dosya uzantisi

        Returns:
            {"text": str, "tables": list, "type": str}
        """
        extension = extension.lower().lstrip('.')

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Desteklenmeyen dosya formati: {extension}")

        processors = {
            'pdf': self._process_pdf,
            'csv': self._process_csv,
            'xlsx': self._process_excel,
            'xls': self._process_excel,
            'txt': self._process_text,
            'docx': self._process_docx,
            'json': self._process_json
        }

        processor = processors.get(extension, self._process_text)
        return processor(content)

    def _process_pdf(self, content: bytes) -> dict:
        """PDF dosyasini isler"""
        from .pdf_service import PDFService
        pdf_service = PDFService()
        result = pdf_service.extract_text(content)
        result['type'] = 'pdf'
        return result

    def _process_csv(self, content: bytes) -> dict:
        """CSV dosyasini isler"""
        df = pd.read_csv(io.BytesIO(content))
        return {
            "text": df.to_string(index=False),
            "tables": [df.values.tolist()],
            "columns": df.columns.tolist(),
            "row_count": len(df),
            "type": "csv"
        }

    def _process_excel(self, content: bytes) -> dict:
        """Excel dosyasini isler"""
        df = pd.read_excel(io.BytesIO(content))
        return {
            "text": df.to_string(index=False),
            "tables": [df.values.tolist()],
            "columns": df.columns.tolist(),
            "row_count": len(df),
            "type": "excel"
        }

    def _process_text(self, content: bytes) -> dict:
        """Duz metin dosyasini isler"""
        text = content.decode('utf-8', errors='ignore')
        return {
            "text": text,
            "tables": [],
            "type": "text"
        }

    def _process_docx(self, content: bytes) -> dict:
        """Word dosyasini isler"""
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]
            text = '\n'.join(paragraphs)

            tables = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    table_data.append([cell.text for cell in row.cells])
                tables.append(table_data)

            return {
                "text": text,
                "tables": tables,
                "type": "docx"
            }
        except ImportError:
            logger.error("python-docx yuklenmemis")
            raise

    def _process_json(self, content: bytes) -> dict:
        """JSON dosyasini isler"""
        import json
        data = json.loads(content.decode('utf-8'))

        if isinstance(data, list):
            df = pd.DataFrame(data)
            text = df.to_string(index=False)
            tables = [df.values.tolist()]
        else:
            text = json.dumps(data, ensure_ascii=False, indent=2)
            tables = []

        return {
            "text": text,
            "tables": tables,
            "type": "json"
        }

    def get_supported_extensions(self) -> list:
        """Desteklenen uzantilari dondurur"""
        return list(self.SUPPORTED_EXTENSIONS)
